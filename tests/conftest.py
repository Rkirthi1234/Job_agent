"""Shared pytest fixtures: an isolated test DB, temp upload dir, and
dummy resume files generated on the fly (no static fixture files needed).
"""
import pytest
from docx import Document
from fpdf import FPDF
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.config import get_settings
from app.main import app
from app.models.database import Base, get_db

DUMMY_RESUME_TEXT = """John Doe
Email: john.doe@example.com
Phone: +1-555-0100
Location: Austin, TX

Summary
Backend engineer with 5 years of experience building Python APIs and cloud data pipelines.

Skills
Python, FastAPI, SQL, Docker

Experience
Software Engineer, Acme Corp, Jan 2021 - Present
- Built REST APIs used by 10 internal teams
- Migrated batch jobs to Azure Functions

Education
B.S. Computer Science, University of Texas, 2016 - 2020
"""


@pytest.fixture(autouse=True)
def _wellfound_login_isolated_from_dotenv(monkeypatch):
    """Pin Wellfound automatic login OFF, with no credentials, for EVERY test.

    A developer's real .env may have WELLFOUND_AUTO_LOGIN=true and a real
    WELLFOUND_EMAIL/WELLFOUND_PASSWORD. Environment variables outrank the
    .env file in pydantic-settings, so pinning them here guarantees no test
    can ever try to sign in to (or type real credentials into) anything.
    Tests that exercise automatic login (tests/test_wellfound_auth.py) set
    their own FAKE values on top of this."""
    monkeypatch.setenv("WELLFOUND_AUTO_LOGIN", "false")
    monkeypatch.setenv("WELLFOUND_EMAIL", "")
    monkeypatch.setenv("WELLFOUND_PASSWORD", "")
    get_settings.cache_clear()
    yield
    get_settings.cache_clear()


@pytest.fixture()
def test_db(tmp_path):
    """Point the app's DB dependency at a throwaway SQLite file for one test."""
    db_path = tmp_path / "test.db"
    engine = create_engine(f"sqlite:///{db_path}", connect_args={"check_same_thread": False})
    testing_session_local = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    Base.metadata.create_all(bind=engine)

    def override_get_db():
        db = testing_session_local()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    yield
    app.dependency_overrides.clear()


@pytest.fixture()
def client(test_db, tmp_path, monkeypatch):
    """A TestClient with uploads redirected to a temp directory.

    Explicitly forces REAL_APPLICATION_ENABLED=false: this fixture is
    shared by the whole suite (Phase 1-4 tests included) and most of
    its application-related tests exercise generic validation (not
    found / duplicate / ineligible / source mismatch / adapter-failure
    simulation) that has nothing to do with which adapter is selected.
    Since real_application_enabled now defaults to True in production
    (see app/config.py), pinning it False here keeps those tests
    exercising exactly what they were written to test. Tests that
    specifically need the real adapter or the true production default
    use the real_client / production_client fixtures in
    tests/test_applications.py instead.
    """
    monkeypatch.setenv("UPLOAD_DIR", str(tmp_path / "uploads"))
    monkeypatch.setenv("REAL_APPLICATION_ENABLED", "false")
    get_settings.cache_clear()
    yield TestClient(app)
    get_settings.cache_clear()


@pytest.fixture()
def dummy_pdf_path(tmp_path):
    """Generate a minimal but real PDF containing resume-like text."""
    path = tmp_path / "resume.pdf"
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for line in DUMMY_RESUME_TEXT.splitlines():
        pdf.cell(0, 8, txt=line, ln=1)
    pdf.output(str(path))
    return path


@pytest.fixture()
def dummy_docx_path(tmp_path):
    """Generate a minimal but real DOCX containing resume-like text."""
    path = tmp_path / "resume.docx"
    document = Document()
    for line in DUMMY_RESUME_TEXT.splitlines():
        document.add_paragraph(line)
    document.save(str(path))
    return path


@pytest.fixture()
def empty_pdf_path(tmp_path):
    """A structurally valid PDF with no text on any page."""
    path = tmp_path / "empty.pdf"
    pdf = FPDF()
    pdf.add_page()
    pdf.output(str(path))
    return path
